import json
from concurrent.futures import ThreadPoolExecutor
from contextlib import nullcontext
from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock
from uuid import UUID

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pydantic import ValidationError

from backend.dependencies import agent_access, auth
from backend.routers import slides
from backend.schemas.slides import SlideComponentNote, SlideDeck, SlideSpec
from utils import course_materials
from utils.slide_export import PPTX_MEDIA_TYPE, SlideLayoutError, _font_size, presentation_filename, render_presentation


def sample_deck():
    return {
        "title": "Electric circuits",
        "subtitle": "Predict, investigate and explain",
        "theme": "academic",
        "slides": [
            {"layout": "title", "title": "Electric circuits", "subtitle": "A classroom introduction"},
            {"layout": "section", "title": "Series and parallel", "subtitle": "One path or multiple paths"},
            {
                "layout": "content", "title": "A complete circuit",
                "bullets": ["A source provides a potential difference.", "Current flows around a closed path."],
                "speaker_notes": "Ask learners to identify the break in the circuit.",
                "sources": [{"title": "Course notes, section 1", "url": "https://example.com/course/notes"}],
            },
            {
                "layout": "two_column", "title": "Compare the connections",
                "columns": [
                    {"heading": "Series", "bullets": ["One current path.", "The same current passes through each component."]},
                    {"heading": "Parallel", "bullets": ["Multiple current paths.", "Branches share the same potential difference."]},
                ],
            },
            {"layout": "question", "title": "Make a prediction", "bullets": ["What changes when one series lamp is disconnected?"]},
            {"layout": "summary", "title": "Explain your reasoning", "bullets": ["Identify the paths.", "Check the connections.", "Test your prediction."]},
        ],
    }


def scripted_deck():
    data = sample_deck()
    data["slides"][2].update(
        speaker_notes="  First consider an unbroken circuit.\r\nWe will connect this idea to current flow.  ",
        component_notes=[
            {"target": "bullet-2", "text": "Opening any point interrupts the path.\nThe same idea applies to a 电路. 🚀"},
            {"target": "title", "text": "\tA complete circuit provides a continuous path for charge.  "},
            {"target": "bullet-1", "text": "  The source transfers energy: ΔV < 5 & I > 0 are example notation.\rExplain each symbol. "},
        ],
        sources=[
            {"title": "Course notes, section 1", "url": "https://example.com/course/notes"},
            {"title": "Course glossary"},
        ],
    )
    return data


def presentation_text(presentation):
    return "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)


@pytest.mark.parametrize("theme", ["academic", "midnight", "warm"])
def test_pptx_contains_editable_slides_notes_citations_and_bounded_shapes(theme):
    data = sample_deck()
    data["theme"] = theme
    deck = SlideDeck.model_validate(data)
    content = render_presentation(deck)
    assert content.startswith(b"PK")
    assert len(content) < 1_000_000
    presentation = Presentation(BytesIO(content))
    assert len(presentation.slides) == len(deck.slides)
    assert presentation.slide_width / presentation.slide_height == pytest.approx(16 / 9, abs=0.001)
    assert presentation.core_properties.title == deck.title
    for slide, spec in zip(presentation.slides, deck.slides):
        texts = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert spec.title in texts
        assert spec.subtitle in texts
        for bullet in spec.bullets:
            assert bullet in texts
        for column in spec.columns:
            assert column.heading in texts
            assert all(bullet in texts for bullet in column.bullets)
        assert spec.speaker_notes in slide.notes_slide.notes_text_frame.text
        for source in spec.sources:
            assert source.title in texts
            assert source.url in slide.notes_slide.notes_text_frame.text
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= presentation.slide_width + 2
            assert shape.top + shape.height <= presentation.slide_height + 2
    citation = next(shape for shape in presentation.slides[2].shapes if shape.has_text_frame and "Course notes" in shape.text)
    assert citation.text_frame.paragraphs[0].runs[0].hyperlink.address == "https://example.com/course/notes"


def test_unicode_content_and_twenty_slides_survive_export():
    data = sample_deck()
    data["slides"] = [{
        "layout": "content", "title": f"Lesson {number + 1}",
        "bullets": ["Voltage \u0394V and current \u0399", "\u7535\u8def"],
        "speaker_notes": "Explain the notation before calculating.",
    } for number in range(20)]
    presentation = Presentation(BytesIO(render_presentation(SlideDeck.model_validate(data))))
    assert len(presentation.slides) == 20
    assert "\u0394V" in presentation_text(presentation)
    assert "\u7535\u8def" in presentation_text(presentation)


def test_long_source_labels_keep_full_reference_on_slide_and_in_notes():
    data = sample_deck()
    source_title = "W" * 120
    data["slides"][2]["sources"][0]["title"] = source_title
    presentation = Presentation(BytesIO(render_presentation(SlideDeck.model_validate(data))))
    slide = presentation.slides[2]
    assert source_title in slide.notes_slide.notes_text_frame.text
    label = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text.startswith("[1]"))
    assert label.text == f"[1] {source_title}"
    assert label.text_frame.paragraphs[0].runs[0].hyperlink.address == "https://example.com/course/notes"


def test_legacy_decks_default_to_independent_empty_component_notes_without_inventing_scripts():
    data = sample_deck()
    original = deepcopy(data)
    first = SlideDeck.model_validate(data)
    second = SlideDeck.model_validate(data)
    assert all(slide.component_notes == [] for slide in first.slides)
    presentation = Presentation(BytesIO(render_presentation(first)))
    assert presentation.slides[0].notes_slide.notes_text_frame.text == ""
    assert presentation.slides[2].notes_slide.notes_text_frame.text == (
        "Ask learners to identify the break in the circuit.\n\n"
        "Sources:\n[1] Course notes, section 1 - https://example.com/course/notes"
    )
    first.slides[0].component_notes.append(SlideComponentNote(target="title", text="Introduce the topic."))
    assert first.slides[1].component_notes == second.slides[0].component_notes == []
    assert data == original


@pytest.mark.parametrize("include_intro", [False, True])
@pytest.mark.parametrize("include_scripts", [False, True])
@pytest.mark.parametrize("include_sources", [False, True])
def test_native_notes_preserve_exact_stored_scripts_intro_labels_and_sources(
    include_intro, include_scripts, include_sources,
):
    data = scripted_deck()
    spec = data["slides"][2]
    if not include_intro:
        spec["speaker_notes"] = ""
    if not include_scripts:
        spec["component_notes"] = []
    if not include_sources:
        spec["sources"] = []
    original = deepcopy(data)
    deck = SlideDeck.model_validate(data)
    assert deck.slides[2].speaker_notes == spec["speaker_notes"]
    assert [note.model_dump() for note in deck.slides[2].component_notes] == spec["component_notes"]
    expected_sections = []
    if include_intro:
        expected_sections.append(spec["speaker_notes"])
    if include_scripts:
        expected_sections.append(
            "Component scripts:\n"
            "bullet-2:\nOpening any point interrupts the path.\nThe same idea applies to a 电路. 🚀\n\n"
            "title:\n\tA complete circuit provides a continuous path for charge.  \n\n"
            "bullet-1:\n  The source transfers energy: ΔV < 5 & I > 0 are example notation.\rExplain each symbol. "
        )
    if include_sources:
        expected_sections.append(
            "Sources:\n[1] Course notes, section 1 - https://example.com/course/notes\n[2] Course glossary"
        )
    presentation = Presentation(BytesIO(render_presentation(deck)))
    expected = "\n\n".join(expected_sections)
    assert presentation.slides[2].notes_slide.notes_text_frame.text == expected
    saved = BytesIO()
    presentation.save(saved)
    assert Presentation(BytesIO(saved.getvalue())).slides[2].notes_slide.notes_text_frame.text == expected
    for note in spec["component_notes"]:
        assert note["text"] not in presentation_text(presentation)
    assert data == original


@pytest.mark.parametrize("slide_index", range(6))
def test_component_targets_match_existing_components_on_every_legacy_layout(slide_index):
    data = sample_deck()["slides"][slide_index]
    targets = ["title"]
    if data.get("subtitle"):
        targets.append("subtitle")
    targets.extend(f"bullet-{index}" for index in range(1, len(data.get("bullets", [])) + 1))
    for column_index, column in enumerate(data.get("columns", []), start=1):
        targets.append(f"column-{column_index}")
        targets.extend(f"column-{column_index}-bullet-{index}" for index in range(1, len(column["bullets"]) + 1))
    data["component_notes"] = [{"target": target, "text": f"Explanation for {target}."} for target in targets]
    spec = SlideSpec.model_validate(data)
    assert [note.target for note in spec.component_notes] == targets


@pytest.mark.parametrize("note", [
    {}, {"target": "title"}, {"text": "Explanation."},
    {"target": "", "text": "Explanation."},
    {"target": " ", "text": "Explanation."},
    {"target": None, "text": "Explanation."},
    {"target": 1, "text": "Explanation."},
    {"target": " title", "text": "Explanation."},
    {"target": "Title", "text": "Explanation."},
    {"target": "bullet-01", "text": "Explanation."},
    {"target": "bullet-0", "text": "Explanation."},
    {"target": "bullet-6", "text": "Explanation."},
    {"target": "column-0", "text": "Explanation."},
    {"target": "column-3", "text": "Explanation."},
    {"target": "column-1-bullet-0", "text": "Explanation."},
    {"target": "column-2-bullet-5", "text": "Explanation."},
    {"target": "source-1", "text": "Explanation."},
    {"target": "title", "text": ""},
    {"target": "title", "text": " \t\r\n\u2003 "},
    {"target": "title", "text": None},
    {"target": "title", "text": 123},
    {"target": "title", "text": "Explanation.", "unexpected": True},
])
def test_component_notes_reject_missing_invalid_empty_or_extra_fields(note):
    data = sample_deck()["slides"][2]
    data["component_notes"] = [note]
    with pytest.raises(ValidationError):
        SlideSpec.model_validate(data)


@pytest.mark.parametrize(("slide_index", "target"), [
    (0, "bullet-1"), (1, "bullet-1"), (2, "subtitle"), (2, "bullet-3"),
    (2, "column-1"), (2, "column-1-bullet-1"), (3, "bullet-1"),
    (3, "column-1-bullet-3"), (3, "column-2-bullet-4"),
])
def test_component_notes_cannot_target_absent_or_out_of_bounds_components(slide_index, target):
    data = sample_deck()["slides"][slide_index]
    data["component_notes"] = [{"target": target, "text": "Explanation."}]
    with pytest.raises(ValidationError, match="existing slide components"):
        SlideSpec.model_validate(data)


def test_component_notes_reject_blank_subtitle_and_duplicate_targets():
    data = sample_deck()["slides"][2]
    data.update(subtitle=" \t ", component_notes=[{"target": "subtitle", "text": "Explanation."}])
    with pytest.raises(ValidationError, match="existing slide components"):
        SlideSpec.model_validate(data)
    data["component_notes"] = [
        {"target": "bullet-1", "text": "First explanation."},
        {"target": "bullet-1", "text": "A different explanation."},
    ]
    with pytest.raises(ValidationError, match="unique"):
        SlideSpec.model_validate(data)


@pytest.mark.parametrize("notes", [None, {}, "Not an array", [None], [1]])
def test_component_notes_reject_non_array_or_non_object_values(notes):
    data = sample_deck()["slides"][2]
    data["component_notes"] = notes
    with pytest.raises(ValidationError):
        SlideSpec.model_validate(data)


def test_twelve_component_notes_are_allowed_but_thirteen_are_not():
    data = sample_deck()["slides"][3]
    data["subtitle"] = "Compare both sections"
    for column in data["columns"]:
        column["bullets"] = ["One", "Two", "Three", "Four"]
    targets = ["title", "subtitle", "column-1", "column-2", *[
        f"column-{column}-bullet-{bullet}" for column in (1, 2) for bullet in range(1, 5)
    ]]
    data["component_notes"] = [{"target": target, "text": "A concrete explanation."} for target in targets]
    assert len(SlideSpec.model_validate(data).component_notes) == 12
    data["component_notes"].append({"target": "title", "text": "One too many."})
    with pytest.raises(ValidationError, match="at most 12"):
        SlideSpec.model_validate(data)


@pytest.mark.parametrize("text", ["🚀" * 1200, "e\u0301" * 600, " " * 1199 + "x"])
def test_component_text_limit_counts_unicode_code_points_without_trimming(text):
    note = SlideComponentNote(target="title", text=text)
    assert note.text == text and len(note.text) == 1200
    with pytest.raises(ValidationError):
        SlideComponentNote(target="title", text=text + "x")


@pytest.mark.parametrize("character", [
    "\x00", "\x01", "\x08", "\x0b", "\x0c", "\x1f", "\ud800", "\udfff", "\ufffe", "\uffff",
])
@pytest.mark.parametrize("field", ["speaker_notes", "component_notes"])
def test_all_script_fields_reject_xml_unsafe_characters(character, field):
    data = sample_deck()["slides"][2]
    text = f"Explain {character} safely."
    data[field] = text if field == "speaker_notes" else [{"target": "title", "text": text}]
    with pytest.raises(ValidationError):
        SlideSpec.model_validate(data)


@pytest.mark.parametrize("overflow_field", ["speaker_notes", "component_notes"])
def test_combined_script_limit_is_8000_unicode_code_points(overflow_field):
    data = sample_deck()["slides"][2]
    data["bullets"] = ["One", "Two", "Three", "Four"]
    data["speaker_notes"] = "🚀" * 3999
    data["component_notes"] = [
        {"target": f"bullet-{index + 1}", "text": "电" * length}
        for index, length in enumerate([1200, 1200, 1200, 401])
    ]
    spec = SlideSpec.model_validate(data)
    assert len(spec.speaker_notes) + sum(len(note.text) for note in spec.component_notes) == 8000
    if overflow_field == "speaker_notes":
        data["speaker_notes"] += "x"
    else:
        data["component_notes"][-1]["text"] += "x"
    with pytest.raises(ValidationError, match="8000 Unicode code points"):
        SlideSpec.model_validate(data)


@pytest.mark.parametrize(("layout", "bullets"), [
    ("process", ["Choose a source.", "Close the loop."]),
    ("process", ["Choose.", "Connect.", "Check.", "Measure.", "Explain."]),
    ("timeline", ["Initial state.", "Final state."]),
    ("timeline", ["Start.", "Observe.", "Measure.", "Compare.", "Conclude."]),
    ("quote", ["A complete path is necessary for current."]),
    ("key_stat", ["3 paths"]),
])
@pytest.mark.parametrize("theme", ["academic", "midnight", "warm"])
def test_new_layouts_preserve_native_editable_shapes_bounded_content_and_component_scripts(layout, bullets, theme):
    data = {
        "title": "Circuit lesson", "theme": theme,
        "slides": [{
            "layout": layout, "title": "Circuit lesson", "subtitle": "Course notes: context and attribution",
            "bullets": bullets, "speaker_notes": "Introduce the next idea.",
            "component_notes": [
                {"target": f"bullet-{index + 1}", "text": f"Detailed explanation for item {index + 1}."}
                for index in range(len(bullets))
            ] + [{"target": "subtitle", "text": "Explain the source or caption."}],
            "sources": [{"title": "Course notes", "url": "https://example.com/course/notes"}],
        }],
    }
    deck = SlideDeck.model_validate(data)
    presentation = Presentation(BytesIO(render_presentation(deck)))
    slide = presentation.slides[0]
    texts = presentation_text(presentation)
    assert all(text in texts for text in [*bullets, data["slides"][0]["subtitle"]])
    for shape in slide.shapes:
        assert shape.shape_type != MSO_SHAPE_TYPE.PICTURE
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= presentation.slide_width + 2
        assert shape.top + shape.height <= presentation.slide_height + 2
    if layout in {"process", "timeline"}:
        markers = [
            shape for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.OVAL
        ]
        assert len(markers) == len(bullets)
        assert sum(shape.shape_type == MSO_SHAPE_TYPE.LINE for shape in slide.shapes) == (
            len(bullets) - 1 if layout == "timeline" else 0
        )
    for note in deck.slides[0].component_notes:
        assert f"{note.target}:\n{note.text}" in slide.notes_slide.notes_text_frame.text


@pytest.mark.parametrize(("layout", "bullets"), [
    ("process", []), ("process", ["One"]), ("process", ["Step"] * 6),
    ("timeline", []), ("timeline", ["One"]), ("timeline", ["Event"] * 6),
    ("quote", []), ("quote", ["One", "Two"]), ("quote", ["x" * 161]),
    ("key_stat", []), ("key_stat", ["One", "Two"]), ("key_stat", ["🚀" * 41]),
])
def test_new_layouts_reject_wrong_shapes_and_text_bounds(layout, bullets):
    with pytest.raises(ValidationError):
        SlideSpec.model_validate({"layout": layout, "title": "Lesson", "bullets": bullets})


@pytest.mark.parametrize("layout", ["process", "timeline", "quote", "key_stat"])
def test_new_layouts_reject_columns(layout):
    data = {
        "layout": layout, "title": "Lesson",
        "bullets": ["One", "Two"] if layout in {"process", "timeline"} else ["One"],
        "columns": sample_deck()["slides"][3]["columns"],
    }
    with pytest.raises(ValidationError, match="Only two-column"):
        SlideSpec.model_validate(data)


def test_quote_and_statistic_accept_exact_unicode_text_limits():
    assert len(SlideSpec(layout="quote", title="Quote", bullets=["电" * 160]).bullets[0]) == 160
    assert len(SlideSpec(layout="key_stat", title="Statistic", bullets=["🚀" * 40]).bullets[0]) == 40


def test_dense_content_fails_instead_of_truncating_or_producing_unreadable_text():
    with pytest.raises(SlideLayoutError, match="too dense"):
        _font_size(["W" * 160] * 5, 2, 1, 24, 16)
    data = sample_deck()
    data["slides"][3]["subtitle"] = "W" * 180
    data["slides"][3]["columns"][0]["bullets"] = ["W" * 140] * 4
    with pytest.raises(SlideLayoutError, match="too dense"):
        render_presentation(SlideDeck.model_validate(data))


@pytest.mark.parametrize("mutation", [
    lambda data: data.update(slides=[]),
    lambda data: data.update(slides=data["slides"] * 4),
    lambda data: data.update(title=" "),
    lambda data: data.update(filename="../../someone-else.pptx"),
    lambda data: data["slides"][0].update(layout="html"),
    lambda data: data["slides"][0].update(bullets=["Not allowed in this layout"]),
    lambda data: data["slides"][2].update(bullets=[]),
    lambda data: data["slides"][2].update(bullets=["x" * 161]),
    lambda data: data["slides"][2].update(speaker_notes="x" * 4001),
    lambda data: data["slides"][2].update(title="Unsupported\x00character"),
    lambda data: data["slides"][3].update(columns=[]),
    lambda data: data["slides"][3].update(bullets=["Not allowed with columns"]),
])
def test_slide_spec_rejects_invalid_or_unbounded_payloads(mutation):
    data = sample_deck()
    mutation(data)
    with pytest.raises(ValidationError):
        SlideDeck.model_validate(data)


@pytest.mark.parametrize("url", [
    "javascript:alert(1)", "file:///private/materials", "https://user:password@example.com",
    "https://example.com/\nprivate", "https:\\\\example.com", "https://",
])
def test_source_urls_reject_unsafe_schemes_credentials_and_controls(url):
    data = sample_deck()
    data["slides"][2]["sources"][0]["url"] = url
    with pytest.raises(ValidationError):
        SlideDeck.model_validate(data)


def test_parallel_decks_do_not_share_content_ids_or_files(tmp_path, monkeypatch):
    from agent_tools.custom.add_slides import AddSlidesTool

    monkeypatch.chdir(tmp_path)
    first = sample_deck()
    second = sample_deck()
    first["title"] = first["slides"][0]["title"] = "Deck alpha"
    second["title"] = second["slides"][0]["title"] = "Deck beta"
    tool = AddSlidesTool()
    with ThreadPoolExecutor(max_workers=2) as pool:
        blocks = list(pool.map(tool.execute, [first, second]))
    assert blocks[0]["slidesId"] != blocks[1]["slidesId"]
    for block, own_title, other_title in zip(blocks, ["Deck alpha", "Deck beta"], ["Deck beta", "Deck alpha"]):
        UUID(block["slidesId"])
        presentation = Presentation(BytesIO(render_presentation(SlideDeck.model_validate(block["deck"]))))
        text = presentation_text(presentation)
        assert own_title in text
        assert other_title not in text
        assert json.loads(tool.output(block, {}))["slide_count"] == 6
    assert list(tmp_path.iterdir()) == []


def test_schema_and_tool_registration_match():
    from agent_tools.custom.add_slides import ADD_SLIDES_TOOL_DEFINITION
    from agent_tools.custom.declare_plan import DeclarePlanTool
    from azure_services.agents.agent_creation import AgentToolBuilder

    parameters = ADD_SLIDES_TOOL_DEFINITION["parameters"]
    schema = SlideDeck.model_json_schema()
    assert parameters["properties"]["slides"]["maxItems"] == schema["properties"]["slides"]["maxItems"]
    assert parameters["properties"]["slides"]["minItems"] == schema["properties"]["slides"]["minItems"]
    assert parameters["properties"]["theme"]["enum"] == schema["properties"]["theme"]["enum"]
    assert parameters["additionalProperties"] is False
    assert DeclarePlanTool().execute({"tools": ["add_message", "add_slides"]})["status"] == "accepted"
    builder = AgentToolBuilder(agent_name="course-example")
    builder._add_function_tools()
    assert sum(tool.name == "add_slides" for tool in builder._tools) == 1


def test_tool_dispatch_and_raw_function_fallback_emit_slides_not_document():
    from base_agents.general_agent import BLOCK_START_EVENTS, GeneralAgent, _try_parse_raw_function_call

    agent = GeneralAgent.__new__(GeneralAgent)
    agent.agent_name = "course-example"
    data = sample_deck()
    result = agent._dispatch_tool_call("add_slides", json.dumps(data), "call-1", "conversation-1", "user-1")
    assert BLOCK_START_EVENTS["add_slides"] == "slides_start"
    assert [event[0] for event in result["yield_events"]] == ["slides"]
    assert json.loads(result["output"])["status"] == "created"
    assert _try_parse_raw_function_call(json.dumps({"name": "add_slides", "arguments": data})) == [("add_slides", data)]
    events = list(agent._handle_plain_text_fallback("conversation-1", json.dumps({"name": "add_slides", "arguments": data})))
    assert "slides_start" in [event[0] for event in events]
    assert "slides" in [event[0] for event in events]
    assert "document" not in [event[0] for event in events]
    rejected = agent._dispatch_tool_call("add_slides", '{"title":"Empty","slides":[]}', "call-2", "conversation-1", "user-1")
    assert rejected["output"].startswith("Error: ")
    assert [event[0] for event in rejected["yield_events"]] == ["block_cancel"]


def test_agui_preserves_the_deck_and_both_script_fields():
    from agent_tools.custom.add_slides import AddSlidesTool
    from backend.agui import AGUITranslator

    data = sample_deck()
    data["slides"][2]["component_notes"] = [
        {"target": "bullet-1", "text": "Charges gain energy from the source."},
    ]
    block = AddSlidesTool().execute(data)
    frames = list(AGUITranslator().run([
        ("slides_start", "", "conversation-1"),
        ("slides", json.dumps(block), "conversation-1"),
        ("done", "", "conversation-1"),
    ]))
    assert frames[-1]["type"] == "RUN_FINISHED"
    encoded = json.dumps(frames)
    assert block["slidesId"] in encoded
    assert "Slides" in encoded
    assert block["deck"]["slides"][2]["speaker_notes"] in encoded
    assert block["deck"]["slides"][2]["component_notes"][0]["text"] in encoded


@pytest.fixture
def export_client(monkeypatch):
    profile = {"id": "teacher-1", "role": "teacher", "status": "active"}
    course = {"id": "course-example", "_etag": "course-revision", "createdById": "teacher-1", "teacherIds": [], "studentIds": ["student-1"]}
    monkeypatch.setattr(auth, "verify_token", lambda token: {"sub": profile["id"]} if token == "test-token" else None)
    monkeypatch.setattr(auth, "load_profile", lambda _user: profile)
    monkeypatch.setattr(agent_access, "student_assignment_ids", lambda _user: [])
    monkeypatch.setattr(course_materials, "load_course", lambda name: course if name == course["id"] else None)
    monkeypatch.setattr(course_materials, "load_setup", lambda _name: {"sessionUuid": "11111111-1111-4111-8111-111111111111"})
    app = FastAPI()
    app.include_router(slides.router)
    return TestClient(app), profile


def test_export_authentication_course_access_and_private_download(export_client):
    client, profile = export_client
    url = "/api/agents/course-example/slides/export"
    body = {"deck": sample_deck()}
    assert client.post(url, json=body).status_code == 401
    headers = {"Authorization": "Bearer test-token"}
    profile["status"] = "disabled"
    assert client.post(url, json=body, headers=headers).status_code == 403
    profile["status"] = "active"
    profile["id"] = "unrelated-user"
    assert client.post(url, json=body, headers=headers).status_code == 403
    profile.update(id="student-1", role="student")
    response = client.post(url, json=body, headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == PPTX_MEDIA_TYPE
    assert response.headers["cache-control"] == "private, no-store"
    assert response.headers["content-disposition"] == 'attachment; filename="Electric-circuits.pptx"'
    assert len(Presentation(BytesIO(response.content)).slides) == 6
    assert client.post("/api/agents/missing-course/slides/export", json=body, headers=headers).status_code == 404


def test_export_reports_invalid_input_and_internal_failure(export_client, monkeypatch):
    client, _profile = export_client
    headers = {"Authorization": "Bearer test-token"}
    url = "/api/agents/course-example/slides/export"
    assert client.post(url, json={"deck": {"title": "Empty", "slides": []}}, headers=headers).status_code == 422
    monkeypatch.setattr(slides, "render_presentation", Mock(side_effect=RuntimeError("private-internal-detail")))
    response = client.post(url, json={"deck": sample_deck()}, headers=headers)
    assert response.status_code == 503
    assert "private-internal-detail" not in response.text
    assert presentation_filename("../../\r\n") == "presentation.pptx"


def test_export_surfaces_readable_layout_errors(export_client, monkeypatch):
    client, _profile = export_client
    monkeypatch.setattr(slides, "render_presentation", Mock(side_effect=SlideLayoutError("Slide text is too dense.")))
    response = client.post(
        "/api/agents/course-example/slides/export",
        json={"deck": sample_deck()}, headers={"Authorization": "Bearer test-token"},
    )
    assert response.status_code == 422
    assert response.json()["detail"] == "Slide text is too dense."


@pytest.fixture
def save_client(export_client, monkeypatch):
    from azure_services.persistence import cosmos_db

    create_asset = Mock(return_value={"id": "saved-asset"})
    monkeypatch.setattr(cosmos_db, "create_asset", create_asset)
    return *export_client, create_asset


@pytest.mark.parametrize(("token", "profile_update", "course_name", "expected"), [
    (None, {}, "course-example", 401),
    ("invalid-token", {}, "course-example", 401),
    ("test-token", {"status": "disabled"}, "course-example", 403),
    ("test-token", {"id": "unrelated-teacher"}, "course-example", 403),
    ("test-token", {"id": "unrelated-student", "role": "student"}, "course-example", 403),
    ("test-token", {}, "missing-course", 404),
])
def test_save_requires_authentication_and_course_access(
    save_client, monkeypatch, token, profile_update, course_name, expected,
):
    client, profile, create_asset = save_client
    profile.update(profile_update)
    preflight = Mock()
    monkeypatch.setattr(slides, "render_presentation", preflight)
    headers = {"Authorization": f"Bearer {token}"} if token else {}
    response = client.post(f"/api/agents/{course_name}/slides/save", json={"deck": sample_deck()}, headers=headers)
    assert response.status_code == expected
    preflight.assert_not_called()
    create_asset.assert_not_called()


@pytest.mark.parametrize(("user_id", "role"), [
    ("teacher-1", "teacher"),
    ("student-1", "student"),
    ("admin-1", "admin"),
])
def test_save_uses_verified_owner_and_originating_agent_after_preflight(save_client, monkeypatch, user_id, role):
    client, profile, create_asset = save_client
    profile.update(id=user_id, role=role)
    preflight = Mock(wraps=render_presentation)
    timeline = Mock()
    timeline.attach_mock(preflight, "render")
    timeline.attach_mock(create_asset, "create")
    monkeypatch.setattr(slides, "render_presentation", preflight)
    deck = SlideDeck.model_validate(sample_deck()).model_dump(mode="json")
    response = client.post(
        "/api/agents/course-example/slides/save?user_id=other-user&isPublic=true&threadId=other-thread",
        json={"deck": deck}, headers={"Authorization": "Bearer test-token"},
    )
    assert response.status_code == 200
    assert response.headers["cache-control"] == "private, no-store"
    saved = response.json()
    assert set(saved) == {"assetId", "block"}
    assert saved["assetId"] == "saved-asset"
    block = saved["block"]
    assert set(block) == {"type", "slidesId", "title", "deck"}
    assert UUID(block["slidesId"]).version == 4
    assert block["type"] == "slides" and block["title"] == deck["title"]
    assert block["deck"] == deck
    content = create_asset.call_args.kwargs["content"]
    assert json.loads(content) == {**block, "agentId": "course-example"}
    create_asset.assert_called_once_with(
        user_id=user_id, title=deck["title"], category="presentation", asset_type="json",
        content=content, agent_id="course-example", is_public=False,
    )
    assert [call[0] for call in timeline.mock_calls] == ["render", "create"]
    assert preflight.call_args.args[0].model_dump(mode="json") == deck


@pytest.mark.parametrize("with_scripts", [False, True])
def test_save_creates_new_private_assets_without_overwriting_original(export_client, monkeypatch, with_scripts):
    from azure_services.persistence import cosmos_db

    client, _profile = export_client
    original_block = {
        "type": "slides", "slidesId": "11111111-1111-4111-8111-111111111111",
        "title": "Electric circuits", "deck": scripted_deck() if with_scripts else sample_deck(),
        "agentId": "course-example",
    }
    original_asset = {
        "id": "source-asset", "userId": "teacher-1", "category": "presentation", "type": "json",
        "content": json.dumps(original_block), "isPublic": False,
        "agentId": "course-example", "threadId": "source-thread", "messageId": "source-message",
    }
    original_message = {"id": "source-message", "role": "assistant", "blocks": [original_block]}
    stored_assets = {original_asset["id"]: deepcopy(original_asset)}
    original_snapshot = deepcopy((original_asset, original_message))

    def create_item(*, body):
        assert body["id"] not in stored_assets
        stored_assets[body["id"]] = deepcopy(body)
        return deepcopy(body)

    container = Mock()
    container.create_item.side_effect = create_item
    monkeypatch.setattr(cosmos_db, "_get_assets_container", lambda: container)
    edited_deck = SlideDeck.model_validate(deepcopy(original_block["deck"])).model_dump(mode="json")
    edited_deck["title"] = edited_deck["slides"][0]["title"] = "My revised circuit lesson"
    edited_deck["slides"][2]["speaker_notes"] = "My private teaching notes."
    edited_deck["slides"][2]["component_notes"] = [
        {"target": "bullet-1", "text": "  My private explanation of how the source provides energy.  "},
    ]
    saved_copies = []
    for _ in range(2):
        response = client.post(
            "/api/agents/course-example/slides/save", json={"deck": edited_deck},
            headers={"Authorization": "Bearer test-token"},
        )
        assert response.status_code == 200
        saved = response.json()
        saved_copies.append(saved)
        record = stored_assets[saved["assetId"]]
        assert UUID(record["id"]).version == 4
        assert record["id"] != original_asset["id"]
        assert saved["block"]["slidesId"] != original_block["slidesId"]
        assert UUID(saved["block"]["slidesId"]).version == 4
        assert record["userId"] == "teacher-1"
        assert record["category"] == "presentation" and record["type"] == "json"
        assert record["title"] == edited_deck["title"]
        assert record["agentId"] == "course-example"
        assert record["isPublic"] is False
        assert record["threadId"] is None and record["messageId"] is None
        assert json.loads(record["content"]) == {**saved["block"], "agentId": "course-example"}
        assert saved["block"]["deck"] == edited_deck
    assert saved_copies[0]["assetId"] != saved_copies[1]["assetId"]
    assert saved_copies[0]["block"]["slidesId"] != saved_copies[1]["block"]["slidesId"]
    assert len(stored_assets) == 3
    assert stored_assets[original_asset["id"]] == original_snapshot[0]
    assert (original_asset, original_message) == original_snapshot
    assert container.create_item.call_count == 2
    container.read_item.assert_not_called()
    container.replace_item.assert_not_called()
    container.upsert_item.assert_not_called()
    container.delete_item.assert_not_called()


@pytest.mark.parametrize(("field", "value"), [
    ("userId", "other-user"), ("user_id", "other-user"), ("ownerId", "other-user"),
    ("isPublic", True), ("visibility", "public"), ("threadId", "source-thread"),
    ("messageId", "source-message"), ("slidesId", "source-slides"), ("assetId", "source-asset"),
    ("agentId", "other-agent"),
])
def test_save_rejects_owner_visibility_and_source_identifiers(save_client, monkeypatch, field, value):
    client, _profile, create_asset = save_client
    preflight = Mock()
    monkeypatch.setattr(slides, "render_presentation", preflight)
    response = client.post(
        "/api/agents/course-example/slides/save", json={"deck": sample_deck(), field: value},
        headers={"Authorization": "Bearer test-token"},
    )
    assert response.status_code == 422
    preflight.assert_not_called()
    create_asset.assert_not_called()


@pytest.mark.parametrize("mutation", [
    lambda body: body.clear(),
    lambda body: body.update(deck={"title": "Empty", "slides": []}),
    lambda body: body["deck"].update(slidesId="source-slides"),
    lambda body: body["deck"]["slides"][2].update(bullets=["x" * 161]),
    lambda body: body["deck"]["slides"][2].update(component_notes=[{"target": "bullet-3", "text": "Missing component."}]),
    lambda body: body["deck"]["slides"][2].update(component_notes=[
        {"target": "title", "text": "Introduction."}, {"target": "title", "text": "Duplicate."},
    ]),
])
def test_save_invalid_payload_does_not_render_or_write(save_client, monkeypatch, mutation):
    client, _profile, create_asset = save_client
    preflight = Mock()
    monkeypatch.setattr(slides, "render_presentation", preflight)
    body = {"deck": sample_deck()}
    mutation(body)
    response = client.post(
        "/api/agents/course-example/slides/save", json=body, headers={"Authorization": "Bearer test-token"},
    )
    assert response.status_code == 422
    preflight.assert_not_called()
    create_asset.assert_not_called()


def test_save_dense_deck_is_readably_rejected_without_writing(save_client):
    client, _profile, create_asset = save_client
    deck = sample_deck()
    deck["slides"][3]["subtitle"] = "W" * 180
    deck["slides"][3]["columns"][0]["bullets"] = ["W" * 140] * 4
    response = client.post(
        "/api/agents/course-example/slides/save", json={"deck": deck}, headers={"Authorization": "Bearer test-token"},
    )
    assert response.status_code == 422
    assert response.json()["detail"] == (
        "Slide text is too dense. Shorten the title or bullets, or split the content into more slides."
    )
    create_asset.assert_not_called()


@pytest.mark.parametrize("error", [
    RuntimeError("private-render-detail"),
    SlideLayoutError("private-layout-detail"),
])
def test_save_preflight_does_not_expose_exception_details_or_write(save_client, monkeypatch, error):
    client, _profile, create_asset = save_client
    monkeypatch.setattr(slides, "render_presentation", Mock(side_effect=error))
    response = client.post(
        "/api/agents/course-example/slides/save", json={"deck": sample_deck()},
        headers={"Authorization": "Bearer test-token"},
    )
    assert response.status_code == (422 if isinstance(error, SlideLayoutError) else 503)
    assert str(error) not in response.text
    create_asset.assert_not_called()


def test_save_course_lookup_failure_is_sanitized_without_writing(save_client, monkeypatch):
    client, _profile, create_asset = save_client
    monkeypatch.setattr(course_materials, "load_course", Mock(side_effect=RuntimeError("private-course-detail")))
    preflight = Mock()
    monkeypatch.setattr(slides, "render_presentation", preflight)
    response = client.post(
        "/api/agents/course-example/slides/save", json={"deck": sample_deck()},
        headers={"Authorization": "Bearer test-token"},
    )
    assert response.status_code == 503
    assert "private-course-detail" not in response.text
    preflight.assert_not_called()
    create_asset.assert_not_called()


@pytest.mark.parametrize("error", [
    RuntimeError("private-storage-detail"),
    HTTPException(status_code=500, detail="private-storage-detail"),
])
def test_save_storage_failure_is_logged_sanitized_and_not_retried(save_client, monkeypatch, caplog, error):
    client, _profile, create_asset = save_client
    create_asset.side_effect = error
    monkeypatch.setattr(slides, "render_presentation", Mock(return_value=b"verified-presentation"))
    response = client.post(
        "/api/agents/course-example/slides/save", json={"deck": sample_deck()},
        headers={"Authorization": "Bearer test-token"},
    )
    assert response.status_code == 503
    assert response.json() == {
        "detail": "The presentation copy could not be saved. Check your assets before saving again.",
    }
    assert "private-storage-detail" not in response.text
    create_asset.assert_called_once()
    assert any(
        record.name == slides.__name__ and record.getMessage() == "Presentation copy save could not be confirmed"
        and record.exc_info
        for record in caplog.records
    )


@pytest.mark.parametrize("result", [None, {}, [], {"id": None}, {"id": ""}, {"id": " "}, {"id": 123}])
def test_save_requires_a_real_asset_result_instead_of_fabricating_success(save_client, monkeypatch, result):
    client, _profile, create_asset = save_client
    create_asset.return_value = result
    monkeypatch.setattr(slides, "render_presentation", Mock(return_value=b"verified-presentation"))
    response = client.post(
        "/api/agents/course-example/slides/save", json={"deck": sample_deck()},
        headers={"Authorization": "Bearer test-token"},
    )
    assert response.status_code == 503
    assert set(response.json()) == {"detail"}
    create_asset.assert_called_once()


def test_existing_agent_tool_enabling_preserves_other_configuration(export_client, monkeypatch):
    from azure.ai.projects.models import FunctionTool, PromptAgentDefinition

    client, profile = export_client
    definition = PromptAgentDefinition(
        model="example-model", instructions="Keep the course instructions.",
        tools=[FunctionTool(name="existing_tool", description="An existing tool", parameters={"type": "object", "properties": {}})],
    )
    current = SimpleNamespace(version="1", definition=definition, metadata={"course": "example"}, description="Course description")
    agent_client = Mock()
    agent_client.agents.get.return_value = SimpleNamespace(versions=SimpleNamespace(latest=current))
    agent_client.agents.create_version.side_effect = lambda **kwargs: SimpleNamespace(version="2", definition=kwargs["definition"])
    monkeypatch.setattr(slides, "project_client", lambda: agent_client)
    monkeypatch.setattr(slides, "tool_update_lock", lambda _name: nullcontext())
    headers = {"Authorization": "Bearer test-token"}
    url = "/api/agents/course-example/slides/tool"
    assert client.get(url, headers=headers).json() == {"enabled": False, "agent_version": "1", "update_available": False}
    assert client.post(url, json={"expected_version": "stale"}, headers=headers).status_code == 409
    agent_client.agents.create_version.assert_not_called()
    result = client.post(url, json={"expected_version": "1"}, headers=headers)
    assert result.status_code == 200
    assert result.json() == {"enabled": True, "agent_version": "2", "update_available": False}
    args = agent_client.agents.create_version.call_args.kwargs
    assert args["definition"].instructions == definition.instructions
    assert args["definition"].model == definition.model
    assert [tool.name for tool in args["definition"].tools] == ["existing_tool", "add_slides"]
    assert [tool.name for tool in definition.tools] == ["existing_tool"]
    assert args["metadata"] == current.metadata and args["description"] == current.description
    assert args["retry_total"] == 0
    current.definition = args["definition"]
    current.version = "2"
    agent_client.agents.create_version.reset_mock()
    assert client.post(url, json={"expected_version": "2"}, headers=headers).status_code == 200
    agent_client.agents.create_version.assert_not_called()
    profile.update(id="student-1", role="student")
    assert client.get(url, headers=headers).status_code == 403
    assert client.post(url, json={"expected_version": "2"}, headers=headers).status_code == 403


@pytest.mark.parametrize("stale_field", ["description", "parameters"])
def test_existing_slide_tool_is_only_refreshed_after_confirmed_owner_post(export_client, monkeypatch, stale_field):
    from azure.ai.projects.models import FunctionTool, PromptAgentDefinition
    from utils import metadata_cache
    from utils.tool_definitions import load_tool_definition

    client, _profile = export_client
    canonical = load_tool_definition("add_slides")
    stale = deepcopy(canonical)
    if stale_field == "description":
        stale["description"] = "Legacy slides without per-component narration."
    else:
        slide_properties = stale["parameters"]["properties"]["slides"]["items"]["properties"]
        slide_properties.pop("component_notes")
        slide_properties["layout"]["enum"] = ["title", "section", "content", "two_column", "question", "summary"]
    definition = PromptAgentDefinition(
        model="example-model", instructions="Preserve the course's teaching instructions.",
        tools=[
            FunctionTool(name="existing_tool", description="Keep this tool.", parameters={"type": "object"}),
            FunctionTool(
                name="add_slides", description=stale["description"], parameters=stale["parameters"], strict=False,
            ),
        ],
    )
    original_definition = deepcopy(definition.as_dict())
    current = SimpleNamespace(version="7", definition=definition, metadata={"course": "example"}, description="Course")
    agent_client = Mock()
    agent_client.agents.get.return_value = SimpleNamespace(versions=SimpleNamespace(latest=current))
    agent_client.agents.create_version.side_effect = lambda **kwargs: SimpleNamespace(version="8", definition=kwargs["definition"])
    monkeypatch.setattr(slides, "project_client", lambda: agent_client)
    lease = Mock(side_effect=lambda _name: nullcontext())
    monkeypatch.setattr(slides, "tool_update_lock", lease)
    invalidate = Mock()
    monkeypatch.setattr(metadata_cache, "invalidate_agent_metadata", invalidate)
    headers = {"Authorization": "Bearer test-token"}
    url = "/api/agents/course-example/slides/tool"
    assert client.get(url, headers=headers).json() == {
        "enabled": True, "agent_version": "7", "update_available": True,
    }
    agent_client.agents.create_version.assert_not_called()
    lease.assert_not_called()
    invalidate.assert_not_called()
    assert client.post(url, json={"expected_version": "6"}, headers=headers).status_code == 409
    agent_client.agents.create_version.assert_not_called()
    response = client.post(url, json={"expected_version": "7"}, headers=headers)
    assert response.status_code == 200
    assert response.json() == {"enabled": True, "agent_version": "8", "update_available": False}
    agent_client.agents.create_version.assert_called_once()
    kwargs = agent_client.agents.create_version.call_args.kwargs
    expected_definition = deepcopy(original_definition)
    expected_definition["tools"][1]["description"] = canonical["description"]
    expected_definition["tools"][1]["parameters"] = canonical["parameters"]
    assert kwargs["definition"].as_dict() == expected_definition
    assert definition.as_dict() == original_definition
    assert kwargs["metadata"] == current.metadata and kwargs["description"] == current.description
    assert kwargs["retry_total"] == 0 and kwargs["read_timeout"] == 30
    invalidate.assert_called_once_with("course-example")
    assert lease.call_count == 2
    current.definition = kwargs["definition"]
    current.version = "8"
    agent_client.agents.create_version.reset_mock()
    invalidate.reset_mock()
    assert client.post(url, json={"expected_version": "8"}, headers=headers).json() == response.json()
    agent_client.agents.create_version.assert_not_called()
    invalidate.assert_not_called()


@pytest.mark.parametrize(("token", "profile_update", "expected"), [
    (None, {}, 401),
    ("invalid-token", {}, 401),
    ("test-token", {"status": "disabled"}, 403),
    ("test-token", {"id": "student-1", "role": "student"}, 403),
    ("test-token", {"id": "another-teacher"}, 403),
])
def test_tool_status_and_refresh_require_active_owner_before_cloud_access(
    export_client, monkeypatch, token, profile_update, expected,
):
    client, profile = export_client
    profile.update(profile_update)
    project = Mock()
    lease = Mock()
    monkeypatch.setattr(slides, "project_client", project)
    monkeypatch.setattr(slides, "tool_update_lock", lease)
    headers = {"Authorization": f"Bearer {token}"} if token else {}
    url = "/api/agents/course-example/slides/tool"
    assert client.get(url, headers=headers).status_code == expected
    assert client.post(url, json={"expected_version": "1"}, headers=headers).status_code == expected
    project.assert_not_called()
    lease.assert_not_called()


@pytest.mark.parametrize(("cloud_status", "expected"), [(None, 503), (409, 409), (412, 409)])
def test_tool_refresh_failure_is_sanitized_and_never_retried(export_client, monkeypatch, cloud_status, expected):
    from azure.ai.projects.models import FunctionTool, PromptAgentDefinition
    from utils import metadata_cache

    client, _profile = export_client
    definition = PromptAgentDefinition(
        model="example-model",
        tools=[FunctionTool(name="add_slides", description="Old contract.", parameters={"type": "object"})],
    )
    original = deepcopy(definition.as_dict())
    current = SimpleNamespace(version="1", definition=definition, metadata={}, description="")
    agent_client = Mock()
    agent_client.agents.get.return_value = SimpleNamespace(versions=SimpleNamespace(latest=current))
    error = RuntimeError("private-remote-detail")
    error.status_code = cloud_status
    agent_client.agents.create_version.side_effect = error
    monkeypatch.setattr(slides, "project_client", lambda: agent_client)
    monkeypatch.setattr(slides, "tool_update_lock", lambda _name: nullcontext())
    invalidate = Mock()
    monkeypatch.setattr(metadata_cache, "invalidate_agent_metadata", invalidate)
    response = client.post(
        "/api/agents/course-example/slides/tool", json={"expected_version": "1"},
        headers={"Authorization": "Bearer test-token"},
    )
    assert response.status_code == expected
    assert "private-remote-detail" not in response.text
    agent_client.agents.create_version.assert_called_once()
    assert agent_client.agents.create_version.call_args.kwargs["retry_total"] == 0
    assert definition.as_dict() == original
    invalidate.assert_not_called()


@pytest.mark.parametrize("conversation_id", [None, "conversation-1"])
def test_sse_endpoint_forwards_presentation_and_completion(monkeypatch, conversation_id):
    monkeypatch.setattr("msal.PublicClientApplication", lambda **_kwargs: object())
    from agent_tools.custom.add_slides import AddSlidesTool
    from backend import main

    block = AddSlidesTool().execute(scripted_deck())

    class FakeAgent:
        def start_chat_stream(self, **_kwargs):
            yield "slides_start", "", "conversation-1"
            yield "slides", json.dumps(block), "conversation-1"
            yield "done", "", "conversation-1"

        continue_chat_stream = start_chat_stream

    monkeypatch.setattr(main, "_load_setup_json", lambda _agent: {})
    monkeypatch.setattr(auth, "verify_token", lambda _token: {"sub": "example-user"})
    monkeypatch.setattr(auth, "load_profile", lambda _uid: {"id": "example-user", "role": "student", "status": "active"})
    monkeypatch.setattr(agent_access, "load_agent", lambda name: {"id": name, "studentIds": ["example-user"]})
    monkeypatch.setattr(main, "get_general_agent", lambda **_kwargs: FakeAgent())
    monkeypatch.setattr(main, "with_suggested_queries", lambda stream, _agent, _text: stream)
    monkeypatch.setattr(main, "_record_inferred_progress", lambda **_kwargs: None)
    response = TestClient(main.app).post("/api/agents/course-example/chat/stream", headers={"Authorization": "Bearer example-session"}, json={
        "text": "Make a presentation about circuits.", "user_id": "example-user",
        "inject_profile": False, "thread_id": conversation_id,
    })
    assert response.status_code == 200
    events = [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: ")]
    assert [event["type"] for event in events] == ["slides_start", "slides", "done"]
    assert events[1]["deck"] == block["deck"]
    assert events[1]["slidesId"] == block["slidesId"]
