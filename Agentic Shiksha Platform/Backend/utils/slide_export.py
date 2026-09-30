from __future__ import annotations

import math
import re
import unicodedata
from io import BytesIO
from typing import TYPE_CHECKING

from backend.schemas.slides import SlideDeck

if TYPE_CHECKING:
    from pptx.slide import Slide


PPTX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
THEMES = {
    "academic": {"background": "F8FAFC", "text": "14213D", "accent": "126E82", "surface": "E8F0F5", "muted": "506176"},
    "midnight": {"background": "111827", "text": "F8FAFC", "accent": "67E8F9", "surface": "1F2937", "muted": "CBD5E1"},
    "warm": {"background": "FFF8ED", "text": "49311F", "accent": "A94C22", "surface": "F5E9D8", "muted": "785B45"},
}


class SlideLayoutError(ValueError):
    pass


def presentation_filename(title: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9_-]+", "-", title).strip("-_")[:80]
    return f"{stem or 'presentation'}.pptx"


def _text_units(text: str) -> float:
    return sum(
        1.05 if unicodedata.east_asian_width(char) in {"W", "F"} or char in "MW@#%&"
        else 0.35 if char.isspace()
        else 0.7
        for char in text
    )


def _wrapped_lines(text: str, capacity: float) -> int:
    total = 0
    for line in text.split("\n"):
        used = 0.0
        count = 1
        for word in line.split():
            length = _text_units(word)
            spacing = 0.35 if used else 0
            if used and used + spacing + length > capacity:
                count += 1
                used = 0
            if length > capacity:
                count += math.ceil(length / capacity) - 1
                used = length % capacity or capacity
            else:
                used += (0.35 if used else 0) + length
        total += count
    return total


def _font_size(paragraphs: list[str], width: float, height: float, maximum: int, minimum: int) -> int:
    # Conservative wrapping leaves room for font substitution in desktop PowerPoint.
    for size in range(maximum, minimum - 1, -1):
        lines = sum(_wrapped_lines(text, width * 72 / size) for text in paragraphs)
        needed = lines * size * 1.2 + max(0, len(paragraphs) - 1) * 8
        if needed <= height * 72:
            return size
    raise SlideLayoutError("Slide text is too dense. Shorten the title or bullets, or split the content into more slides.")


def _text(
    slide: Slide,
    paragraphs: list[str],
    x: float,
    y: float,
    width: float,
    height: float,
    color: str,
    *,
    maximum: int = 22,
    minimum: int = 16,
    bold: bool = False,
    link: str | None = None,
    center: bool = False,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Inches, Pt

    size = _font_size(paragraphs, width, height, maximum, minimum)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, text in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
        paragraph.line_spacing = 1.15
        paragraph.space_after = Pt(8 if index < len(paragraphs) - 1 else 0)
        if link:
            for run in paragraph.runs:
                run.hyperlink.address = link


def _panel(slide: Slide, x: float, y: float, width: float, height: float, color: str) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()


def _steps(
    slide: Slide, captions: list[str], top: float, height: float, colors: dict[str, str], *, timeline: bool,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
    from pptx.util import Inches, Pt

    gap = 0.22
    width = (12.1 - gap * (len(captions) - 1)) / len(captions)
    previous = None
    for index, caption in enumerate(captions):
        left = 0.6 + index * (width + gap)
        if not timeline:
            _panel(slide, left, top, width, height, colors["surface"])
        marker_left = left + (width - 0.6) / 2 if timeline else left + 0.2
        marker = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL, Inches(marker_left), Inches(top + 0.18), Inches(0.6), Inches(0.6),
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = RGBColor.from_string(colors["accent"])
        marker.line.fill.background()
        if timeline and previous is not None:
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                previous.left + previous.width, previous.top + previous.height // 2,
                marker.left, marker.top + marker.height // 2,
            )
            connector.begin_connect(previous, 3)
            connector.end_connect(marker, 1)
            connector.line.color.rgb = RGBColor.from_string(colors["accent"])
            connector.line.width = Pt(2)
        _text(
            slide, [str(index + 1)], marker_left, top + 0.29, 0.6, 0.38,
            colors["background"], maximum=20, minimum=20, bold=True, center=True,
        )
        _text(
            slide, [caption], left + 0.2, top + 1.02, width - 0.4, height - 1.2,
            colors["text"], maximum=22, minimum=16, center=timeline,
        )
        previous = marker


def render_presentation(deck: SlideDeck) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = deck.title
    presentation.core_properties.subject = deck.subtitle
    presentation.core_properties.author = "EKALAIVA"
    colors = THEMES[deck.theme]
    for index, spec in enumerate(deck.slides):
        reference_labels = [f"[{number + 1}] {source.title}" for number, source in enumerate(spec.sources)]
        reference_heights = [
            _wrapped_lines(label, 12 * 72 / 10) * 10 * 1.2 / 72 for label in reference_labels
        ]
        references_height = sum(reference_heights) + max(0, len(spec.sources) - 1) * 0.06
        if references_height > 1.6:
            raise SlideLayoutError("Slide text is too dense. Shorten the source labels or split the content into more slides.")
        reference_top = 7.02 - references_height
        content_bottom = min(6.5, reference_top - 0.18) if spec.sources else 6.5
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(colors["background"])
        _panel(slide, 0, 0, 0.12, 7.5, colors["accent"])
        _text(slide, [spec.layout.replace("_", " ").upper()], 0.6, 0.35, 10.5, 0.25, colors["accent"], maximum=11, minimum=11, bold=True)
        _text(slide, [f"{index + 1} / {len(deck.slides)}"], 11.65, 7.12, 1.1, 0.22, colors["muted"], maximum=10, minimum=10)

        if spec.layout in {"title", "section"}:
            _panel(slide, 0.6, 1.25, 1.2, 0.08, colors["accent"])
            _text(slide, [spec.title], 0.6, 1.7, 12.1, 2.45, colors["text"], maximum=40, minimum=28, bold=True)
            if spec.subtitle:
                _text(slide, [spec.subtitle], 0.65, 4.4, 11.7, min(1.6, content_bottom - 4.4), colors["muted"], maximum=24, minimum=18)
        else:
            _text(slide, [spec.title], 0.6, 0.85, 12.1, 1.25, colors["text"], maximum=30, minimum=20, bold=True)
            top = 2.2
            if spec.subtitle and spec.layout not in {"quote", "key_stat"}:
                _text(slide, [spec.subtitle], 0.65, top, 12, 0.8, colors["muted"], maximum=19, minimum=16)
                top += 0.95
            height = content_bottom - top
            if spec.layout == "two_column":
                for column_index, column in enumerate(spec.columns):
                    left = 0.6 + column_index * 6.2
                    _panel(slide, left, top, 5.9, height, colors["surface"])
                    _text(slide, [column.heading], left + 0.22, top + 0.18, 5.46, 0.8, colors["accent"], maximum=23, minimum=18, bold=True)
                    _text(slide, [f"\u2022 {bullet}" for bullet in column.bullets], left + 0.22, top + 1.1, 5.46, height - 1.25, colors["text"], maximum=20, minimum=16)
            elif spec.layout in {"process", "timeline"}:
                _steps(slide, spec.bullets, top, height, colors, timeline=spec.layout == "timeline")
            elif spec.layout == "quote":
                _panel(slide, 0.6, top, 12.1, height, colors["surface"])
                _text(slide, ["\u201c"], 0.85, top + 0.1, 0.65, 1, colors["accent"], maximum=56, minimum=56, bold=True)
                attribution_height = 1.05 if spec.subtitle else 0
                _text(
                    slide, spec.bullets, 1.65, top + 0.3, 10.55, height - 0.65 - attribution_height,
                    colors["text"], maximum=34, minimum=24,
                )
                if spec.subtitle:
                    _text(slide, [spec.subtitle], 1.65, content_bottom - 1.05, 10.55, 0.85, colors["accent"], maximum=20, minimum=16, bold=True)
            elif spec.layout == "key_stat":
                _panel(slide, 0.6, top, 12.1, height, colors["surface"])
                context_height = 1.05 if spec.subtitle else 0
                _text(
                    slide, spec.bullets, 1, top + 0.3, 11.3, height - 0.6 - context_height,
                    colors["accent"], maximum=76, minimum=36, bold=True, center=True,
                )
                if spec.subtitle:
                    _text(slide, [spec.subtitle], 1.1, content_bottom - 1.05, 11.1, 0.85, colors["text"], maximum=22, minimum=16, center=True)
            else:
                if spec.layout in {"question", "summary"}:
                    _panel(slide, 0.6, top, 12.1, height, colors["surface"])
                _text(slide, [f"\u2022 {bullet}" for bullet in spec.bullets], 0.85, top + 0.2, 11.6, height - 0.4, colors["text"], maximum=24, minimum=16)

        for source, label, reference_height in zip(spec.sources, reference_labels, reference_heights):
            _text(slide, [label], 0.65, reference_top, 12, reference_height, colors["muted"], maximum=10, minimum=10, link=source.url)
            reference_top += reference_height + 0.06
        notes = [spec.speaker_notes] if spec.speaker_notes else []
        if spec.component_notes:
            scripts = "\n\n".join(f"{note.target}:\n{note.text}" for note in spec.component_notes)
            notes.append(f"Component scripts:\n{scripts}")
        if spec.sources:
            references = "\n".join(f"[{number + 1}] {source.title}" + (f" - {source.url}" if source.url else "") for number, source in enumerate(spec.sources))
            notes.append(f"Sources:\n{references}")
        notes_frame = slide.notes_slide.notes_text_frame
        if notes_frame is None:
            raise SlideLayoutError("Speaker notes could not be generated.")
        notes_frame.clear()
        for number, text in enumerate("\n\n".join(notes).split("\n")):
            paragraph = notes_frame.paragraphs[0] if number == 0 else notes_frame.add_paragraph()
            # The public setter escapes legal carriage returns into literal "_x000D_" text.
            paragraph.add_run()._r.t.text = text

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
